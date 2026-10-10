"""导入《企业RAG全环节测试集_600MW.xlsx》测试用例总表为本项目标准数据集。

提炼自该表结构（表头在第 5 行、数据自第 6 行起、空行跳过）：
- PAR/CHK（原文解析/切分）→ parse 基线文件（供 parse_benchmark.py 比对）
- RET/RER/ANS/CIT/REF → 问答 golden 文件（DATASET_SPEC V2 规范键）
同一 SCN 场景的 RET/RER/ANS/CIT 归入同一 paraphrase_group（口语改写组）。

用法：
    python scripts/import_xlsx_cases.py --xlsx 企业RAG全环节测试集_600MW.xlsx \\
        --out-golden datasets/xlsx_golden.jsonl \\
        --out-baseline datasets/xlsx_parse_baseline.jsonl
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import Counter
from typing import Any, Dict, List, Tuple

STAGE_MAP = {
    "原文解析": "parse", "切分": "chunk", "召回": "retrieval",
    "重排": "rerank", "回答": "qa", "引用": "qa", "拒答": "refusal",
}

# xlsx 问题类型 → (question_type, refusal_reason|None)
QTYPE_MAP: Dict[str, Tuple[str, str]] = {
    "表格参数": ("table", ""), "表格比较": ("table", ""), "表格对比": ("table", ""),
    "参数限值": ("parameter", ""), "参数处置": ("parameter", ""),
    "区间判断": ("range", ""), "启停条件": ("condition", ""),
    "步骤顺序": ("procedure", ""), "停运步骤": ("procedure", ""),
    "操作步骤": ("procedure", ""), "事故现象": ("accident_symptom", ""),
    "事故原因": ("accident_cause", ""), "事故处置": ("accident_action", ""),
    "跨段综合": ("multi_document", ""), "禁止性规定": ("procedure", ""),
    "对比型": ("compare", ""), "无答案": ("negative", "no_answer"),
    "含糊问题": ("negative", "ambiguous"), "依据不足": ("negative", "insufficient"),
    "错误前提": ("negative", "false_premise"),
    "版本失效（配置场景）": ("negative", "stale_version"),
    "权限不足（配置场景）": ("negative", "no_permission"),
}

TABLE_RE = re.compile(r"表T\d+")
FORBID_RE = re.compile(r"禁止[:：]")


def split_expectation(exp: str) -> Tuple[List[str], List[str]]:
    """阶段预期结果 → (要点[], 禁止[])。禁止段以'禁止：'引导；要点取末段冒号后按；。切分。"""
    exp = (exp or "").strip()
    if not exp:
        return [], []
    m = FORBID_RE.search(exp)
    forbidden: List[str] = []
    if m:
        forbidden = [p.strip(" 。；") for p in re.split(r"[；。]", exp[m.end():])
                     if p.strip(" 。；")]
        exp = exp[:m.start()]
    if "：" in exp or ":" in exp:
        exp = re.split(r"[:：]", exp)[-1]
    points = [p.strip(" 。；") for p in re.split(r"[；。]", exp) if p.strip(" 。；")]
    return points, forbidden


def strip_version(v: str) -> str:
    return re.split(r"[（(]", (v or "").strip())[0].strip()


def source_hash_from_notes(wb) -> str:
    try:
        ws = wb["填写说明"]
    except KeyError:
        return ""
    for row in ws.iter_rows(values_only=True):
        cells = [(c or "") for c in row]
        hit = next((i for i, c in enumerate(cells) if "SHA-256" in str(c).upper()
                    or "SHA-256" in str(c) or "SHA256" in str(c).upper()), None)
        if hit is not None:
            for c in cells[hit + 1:]:
                if c and len(str(c).strip()) >= 16:
                    return str(c).strip()
    return ""


def read_case_rows(xlsx: str):
    try:
        import openpyxl  # noqa: E402
    except ImportError:
        print("需安装 openpyxl：uv add openpyxl / pip install openpyxl",
              file=sys.stderr)
        sys.exit(2)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["测试用例总表"]
    rows = []
    for idx, r in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        if not r[1]:
            continue
        rows.append((idx, [("" if v is None else str(v).strip()) for v in r[:14]]))
    return wb, rows


def build_golden(scn: str, cid: str, stage_cn: str, qtype_cn: str, question: str,
                 expect: str, evidence: str, pos: str, doc: str, ver: str,
                 verify: str, audit: str, note: str, rowno: int,
                 src_hash: str) -> Dict[str, Any]:
    stage = STAGE_MAP.get(stage_cn, "qa")
    qtype, refusal = QTYPE_MAP.get(qtype_cn, ("fact", ""))
    points, forbidden = split_expectation(expect)
    tables = TABLE_RE.findall(pos or "")
    version = strip_version(ver)
    is_refusal = stage == "refusal"
    prefix = cid.split("-")[0] if "-" in cid else ""
    g: Dict[str, Any] = {
        "question_id": cid,
        "question": question,
        "test_stage": stage,
        "question_type": qtype,
        "paraphrase_group": scn if prefix in ("RET", "RER", "ANS", "CIT") else "",
        "difficulty": "medium",
        "risk_level": "medium",
        "answerable": not is_refusal,
        "expected_answer_points": points if not is_refusal else [],
        "mandatory_points": [],
        "forbidden_claims": forbidden,
        "evidence_text": evidence,
        "chapter_path": pos,
        "table_or_item": "、".join(dict.fromkeys(tables)),
        "expected_document": (doc or "").replace(".docx", ""),
        "expected_chunk": "",
        "acceptable_chunks": [],
        "document_ids": [],
        "chunk_ids": [],
        "golden_answer": "；".join(points) if prefix in ("ANS", "CIT") else "",
        "must_cite": prefix == "CIT",
        "should_refuse": is_refusal,
        "refusal_reason": refusal,
        "requires_order": qtype in ("procedure", "accident_action"),
        "requires_unit": qtype in ("table", "parameter", "range"),
        "attachment_chunk_ids": [],
        "source_version": version,
        "source_hash": src_hash,
        "scenario_id": scn,
        "verify_method": verify,
        "source_ref": f"测试用例总表#R{rowno}",
        "needs_chunk_mapping": prefix in ("RET", "RER"),
        "reviewer": "",
        "review_status": "pending" if "待审核" in audit else "pending",
        "remarks": "；".join([s for s in
                              ["测试集环节：" + stage_cn,
                               ("阶段预期：" + expect) if is_refusal else "",
                               ("备注：" + note) if note else ""] if s]),
    }
    return g


def build_baseline(cid: str, qtype_cn: str, target: str, expect: str,
                   evidence: str, pos: str, stage_cn: str, scn: str,
                   rowno: int) -> Dict[str, Any]:
    kind = "table_cell" if ("表格" in qtype_cn or "表" in (target or "")) else "paragraph"
    return {"record_id": cid,
            "kind": kind,
            "text": evidence,
            "chapter_path": pos,
            "position": {"scenario": scn, "xlsx_stage": stage_cn},
            "input_object": target,
            "expectation": expect,
            "source_ref": f"测试用例总表#R{rowno}"}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out-golden", default="datasets/xlsx_golden.jsonl")
    ap.add_argument("--out-baseline", default="datasets/xlsx_parse_baseline.jsonl")
    args = ap.parse_args()

    wb, rows = read_case_rows(args.xlsx)
    src_hash = source_hash_from_notes(wb)
    golden, baseline = [], []
    stages, qtypes, skipped = Counter(), Counter(), []
    for rowno, r in rows:
        scn, cid, _, stage_cn, qtype_cn = r[0], r[1], r[2], r[3], r[4]
        question, expect, evidence = r[5], r[6], r[7]
        pos, doc, ver, verify, audit, note = r[8], r[9], r[10], r[11], r[12], r[13]
        if stage_cn not in STAGE_MAP:
            skipped.append((cid, f"未知环节 {stage_cn}"))
            continue
        prefix = cid.split("-")[0] if "-" in cid else ""
        if prefix in ("PAR", "CHK"):
            if not evidence:
                skipped.append((cid, "无原文证据"))
                continue
            baseline.append(build_baseline(cid, qtype_cn, question, expect,
                                           evidence, pos, stage_cn, scn, rowno))
        else:
            if not question:
                skipped.append((cid, "无测试问题"))
                continue
            golden.append(build_golden(scn, cid, stage_cn, qtype_cn, question,
                                       expect, evidence, pos, doc, ver,
                                       verify, audit, note, rowno, src_hash))
        stages[stage_cn] += 1
        qtypes[qtype_cn] += 1

    for path, items in ((args.out_golden, golden), (args.out_baseline, baseline)):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            for d in items:
                f.write(json.dumps(d, ensure_ascii=False) + "\n")
    print(f"读入 {len(rows)} 条 → golden {len(golden)} 条 ({args.out_golden}）"
          f"+ parse基线 {len(baseline)} 条 ({args.out_baseline})")
    print(f"环节分布 {dict(stages)}")
    print(f"全部审核态待审核；source_hash={src_hash[:12] + '…' if src_hash else '空'}；"
          f"RET/RER {sum(1 for g in golden if g['needs_chunk_mapping'])} 条待入库映射Chunk")
    if skipped:
        print(f"跳过 {len(skipped)} 条：{skipped[:10]}")


if __name__ == "__main__":
    main()
