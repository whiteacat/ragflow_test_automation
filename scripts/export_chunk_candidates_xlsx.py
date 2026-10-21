"""把检索候选导出成 xlsx 供人工勾选正确答案 chunk，再把勾选结果回收成映射。

两步用法：
    # 1) 导出（对每题跑一次检索，TopN 候选一行一个）
    python scripts/export_chunk_candidates_xlsx.py --config config/test.yaml \\
        --golden datasets/xlsx_golden.jsonl datasets/smoke_small.jsonl \\
        --out chunk_review.xlsx --topn 10
    # 2) 人工在「选中」列填 1（可多选），然后回收并写回 golden
    python scripts/export_chunk_candidates_xlsx.py --apply-xlsx chunk_review.xlsx \\
        --mapping-out mapping.jsonl \\
        --in-golden datasets/xlsx_golden.jsonl --out-golden datasets/xlsx_golden.jsonl
"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

HEADERS = ["来源文件", "题号", "环节", "问题", "证据文本", "现有映射",
           "候选排名", "chunk_id", "文档", "分数", "内容摘要(前300字)", "选中(填1)", "备注"]
WIDTHS = [22, 12, 8, 40, 50, 20, 8, 34, 28, 8, 60, 10, 18]
SELECTED = {"1", "y", "yes", "true", "是", "√"}


def _retrieve(client, ds, ret_cfg, question):
    return client.retrieve(
        question, ds,
        page=int(ret_cfg.get("page", 1)),
        page_size=int(ret_cfg.get("page_size", 30)),
        top_k=int(ret_cfg.get("top_k", 1024)),
        similarity_threshold=float(ret_cfg.get("similarity_threshold", 0.2)),
        vector_similarity_weight=float(
            ret_cfg.get("vector_similarity_weight", 0.3)),
        keyword=bool(ret_cfg.get("keyword", False)),
        highlight=bool(ret_cfg.get("highlight", False)),
        rerank_id=ret_cfg.get("rerank_id"))


def _load_golden(paths):
    out = []
    for p in paths:
        for line in open(p, encoding="utf-8"):
            if line.strip():
                out.append((os.path.basename(p), json.loads(line)))
    return out


def _needs_export(g, include_all):
    if g.get("should_refuse") or g.get("answerable") is False:
        return False
    if include_all:
        return True
    return not (g.get("expected_chunk") or g.get("chunk_ids")
                or g.get("acceptable_chunks"))


def export(paths, cfg_path, out, topn, include_all):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from src.config import load_config
    from src.ragflow_client import RAGFlowClient

    cfg = load_config(cfg_path)
    rag, ret_cfg = cfg["ragflow"], cfg.get("retrieval", {})
    client = RAGFlowClient(rag["base_url"], rag["api_key"],
                           timeout_s=int(rag.get("timeout_s", 60)))
    ds = list(rag.get("dataset_ids") or [])

    items = [(src, g) for src, g in _load_golden(paths) if _needs_export(g, include_all)]
    wb = Workbook()
    ws = wb.active
    ws.title = "候选清单"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.freeze_panes = "A2"

    cache, failed = {}, []
    for i, (src, g) in enumerate(items, 1):
        qid, question = g.get("question_id", ""), g.get("question", "")
        mapped = next((x for x in (g.get("expected_chunk"),
                                   (g.get("chunk_ids") or [""])[0]) if x), "")
        print(f"[{i}/{len(items)}] {qid} 检索中…", flush=True)
        if question in cache:
            chunks, err = cache[question]
        else:
            try:
                chunks, err = _retrieve(client, ds, ret_cfg, question)["chunks"], ""
            except Exception as exc:  # noqa: BLE001 — 单题失败记一行，继续其余题
                chunks, err = [], str(exc)[:200]
            cache[question] = (chunks, err)
        if err:
            failed.append(qid)
            ws.append([src, qid, g.get("test_stage", ""), question,
                       (g.get("evidence_text") or "")[:800], mapped, "", "", "", "",
                       f"检索失败：{err}", "", ""])
            print(f"  ✗ {err[:80]}", flush=True)
            continue
        if not chunks:
            ws.append([src, qid, g.get("test_stage", ""), question,
                       (g.get("evidence_text") or "")[:800], mapped, "", "", "", "",
                       "无召回候选（先查检索/阈值）", "", ""])
            print("  ✗ 无召回候选", flush=True)
            continue
        for c in chunks[:topn]:
            ws.append([src, qid, g.get("test_stage", ""), question,
                       (g.get("evidence_text") or "")[:800], mapped, c.rank,
                       c.chunk_id, c.document_name, round(c.score, 4),
                       (c.content or "")[:300], "", ""])
        print(f"  ✓ {len(chunks[:topn])} 个候选", flush=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    guide = wb.create_sheet("填写说明")
    for line in [
        "1. 在「候选清单」的「选中(填1)」列，给正确答案 chunk 填 1；多个 chunk 共同作答就多填。",
        "2. 每题至少选 1 个；不选则该题保持未映射，检索指标不计分母。",
        "3. 保存后执行：",
        "   python scripts/export_chunk_candidates_xlsx.py --apply-xlsx <本文件> \\",
        "       --mapping-out mapping.jsonl --in-golden <golden> --out-golden <golden>",
        "4. 「现有映射」非空表示该题已有 gold，仅在 --all 导出时出现。",
        "5. 同一问题文本的多个题号（如 ANS-001/CIT-001）共用同一批候选，但需分别勾选。",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 100

    wb.save(out)
    print(f"候选清单 {len(items)} 题 → {out}"
          + (f"；检索失败 {len(failed)} 题：{failed[:10]}" if failed else ""))


def _is_selected(v):
    return str(v).strip().lower() in SELECTED if v is not None else False


def read_selection(xlsx_path):
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["候选清单"]
    head = [str(c.value or "") for c in ws[1]]
    idx = {name: head.index(name) for name in
           ("题号", "候选排名", "chunk_id", "选中(填1)", "备注")}

    picked = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        qid = row[idx["题号"]]
        if not qid or not _is_selected(row[idx["选中(填1)"]]):
            continue
        chunk_id = str(row[idx["chunk_id"]] or "").strip()
        if not chunk_id:
            continue
        rank = row[idx["候选排名"]] or 999
        picked.setdefault(str(qid), []).append((float(rank), chunk_id))
    return {qid: [c for _, c in sorted(v)] for qid, v in picked.items()}


def write_mapping(picked, out):
    with open(out, "w", encoding="utf-8") as f:
        for qid, chunks in picked.items():
            f.write(json.dumps({
                "question_id": qid,
                "expected_chunk": chunks[0],
                "acceptable_chunks": chunks[1:],
            }, ensure_ascii=False) + "\n")
    print(f"勾选 {len(picked)} 题 → {out}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="config/test.yaml")
    ap.add_argument("--golden", nargs="+", default=["datasets/xlsx_golden.jsonl"])
    ap.add_argument("--out", default="chunk_review.xlsx")
    ap.add_argument("--topn", type=int, default=10)
    ap.add_argument("--all", action="store_true", help="连已有映射的题一起导出")
    ap.add_argument("--apply-xlsx", default="", help="回收人工勾选结果")
    ap.add_argument("--mapping-out", default="mapping.jsonl")
    ap.add_argument("--in-golden", default="")
    ap.add_argument("--out-golden", default="")
    args = ap.parse_args()

    if not args.apply_xlsx:
        export(args.golden, args.config, args.out, args.topn, args.all)
        return 0

    picked = read_selection(args.apply_xlsx)
    if not picked:
        print("未发现勾选（「选中(填1)」列没有 1），未生成映射。")
        return 1
    write_mapping(picked, args.mapping_out)
    if args.in_golden and args.out_golden:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        from assist_chunk_mapping import apply as apply_mapping  # noqa: E402
        apply_mapping(args.mapping_out, args.in_golden, args.out_golden)
    else:
        print("未给 --in-golden/--out-golden，仅生成映射文件；"
              "可用 assist_chunk_mapping.py --apply 写回。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
