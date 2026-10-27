"""export_review_xlsx 的离线回归：mock per_question 数据，验证 sheet 结构与内容。"""
import json

import pytest

from scripts.export_review_xlsx import (
    OBJECTIVE_METRICS,
    SUBJECTIVE_METRICS,
    load_per_question,
    sheet_objective,
    sheet_review,
    sheet_rules,
    sheet_summary,
)


def _mock_rows():
    return [
        {
            "question_id": "ANS-001", "scenario_id": "SCN-001", "test_stage": "qa",
            "question_type": "table", "question": "水压试验压力是多少？",
            "golden_answer": "水压试验18.84MPa",
            "expected_answer_points": ["水压试验18.84MPa"],
            "forbidden_claims": ["再热器5.7MPa"],
            "predicted_answer": "水压试验压力为18.84MPa",
            "evidence_snippet": "[1] doc.pdf#c1: 水压试验压力18.84MPa",
            "Top@5": 1, "Top@10": 1, "Top@20": 1,
            "MRR": 1.0, "NDCG@10": 1.0,
            "evidence_recall@5": 1.0, "evidence_recall@10": 1.0,
            "fact_recall_rule": 1.0, "fact_recall_relaxed": 1.0,
            "fact_uncertain_count": 0,
            "abstention_correct": 1, "forbidden_fact_hit": 0,
            "citation_complete_rule": None,
            "trace_ok": True, "run_error": None, "chat_error": None,
            "must_cite": False, "should_refuse": False,
            "answerable": True, "refusal_reason": "",
        },
        {
            "question_id": "CIT-001", "scenario_id": "SCN-001", "test_stage": "qa",
            "question_type": "table", "question": "水压试验压力是多少？",
            "golden_answer": "水压试验18.84MPa",
            "expected_answer_points": ["水压试验18.84MPa"],
            "forbidden_claims": [],
            "predicted_answer": "根据规程，水压试验压力为18.84MPa [ID:0]",
            "evidence_snippet": "[1] doc.pdf#c1: 水压试验压力18.84MPa",
            "Top@5": 1, "Top@10": 1, "Top@20": 1,
            "MRR": 1.0, "NDCG@10": 1.0,
            "evidence_recall@5": 1.0, "evidence_recall@10": 1.0,
            "fact_recall_rule": 1.0, "fact_recall_relaxed": 1.0,
            "fact_uncertain_count": 0,
            "abstention_correct": 1, "forbidden_fact_hit": 0,
            "citation_complete_rule": 1,
            "trace_ok": True, "run_error": None, "chat_error": None,
            "must_cite": True, "should_refuse": False,
            "answerable": True, "refusal_reason": "",
        },
        {
            "question_id": "REF-001", "scenario_id": "SCN-019", "test_stage": "refusal",
            "question_type": "negative", "question": "能用5.7MPa做超压试验吗？",
            "golden_answer": "",
            "expected_answer_points": [],
            "forbidden_claims": [],
            "predicted_answer": "根据规程，超压试验压力为23.55MPa",
            "evidence_snippet": "",
            "Top@5": None, "Top@10": None, "Top@20": None,
            "MRR": None, "NDCG@10": None,
            "evidence_recall@5": None, "evidence_recall@10": None,
            "fact_recall_rule": None, "fact_recall_relaxed": None,
            "fact_uncertain_count": None,
            "abstention_correct": 0, "forbidden_fact_hit": 0,
            "citation_complete_rule": None,
            "trace_ok": True, "run_error": None, "chat_error": None,
            "must_cite": False, "should_refuse": True,
            "answerable": False, "refusal_reason": "false_premise",
        },
        {
            "question_id": "ANS-007", "scenario_id": "SCN-007", "test_stage": "qa",
            "question_type": "table", "question": "超时的题",
            "golden_answer": "",
            "expected_answer_points": [],
            "forbidden_claims": [],
            "predicted_answer": "",
            "evidence_snippet": "",
            "Top@5": None, "MRR": None,
            "fact_recall_rule": None,
            "abstention_correct": None,
            "trace_ok": None,
            "run_error": "retrieval: timeout",
            "chat_error": None,
            "must_cite": False, "should_refuse": False,
            "answerable": True, "refusal_reason": "",
        },
        {
            "question_id": "ANS-010", "scenario_id": "SCN-010", "test_stage": "qa",
            "question_type": "procedure", "question": "停泵步骤是什么？",
            "golden_answer": "先确认停止条件，再解除联锁",
            "expected_answer_points": ["先确认停止条件", "解除备用泵联锁"],
            "forbidden_claims": [],
            "predicted_answer": "先解除联锁，再确认停止条件",
            "evidence_snippet": "[1] doc.pdf#c2: 停泵步骤",
            "Top@5": 0, "Top@10": 1, "Top@20": 1,
            "MRR": 0.5, "NDCG@10": 0.5,
            "evidence_recall@5": 0.5, "evidence_recall@10": 1.0,
            "fact_recall_rule": 0.5, "fact_recall_relaxed": 0.5,
            "fact_uncertain_count": 0,
            "abstention_correct": 1, "forbidden_fact_hit": 0,
            "citation_complete_rule": None,
            "trace_ok": True, "run_error": None, "chat_error": None,
            "must_cite": False, "should_refuse": False,
            "answerable": True, "refusal_reason": "",
        },
    ]


def _mock_summary():
    return {
        "metrics": {
            "RAG-01": {"value": 0.371, "n": 70},
            "GEN-02": {"value": 0.056, "n": 70, "covered": 11, "total": 198},
        },
        "verdicts": {
            "RAG-01": {"pass": False, "desc": "0.371 >= 0.9"},
            "GEN-02": {"pass": False, "desc": "0.056 >= 0.95"},
        },
    }


# ---- Sheet 结构 ----

def test_rules_sheet_has_all_metrics():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_rules(wb)
    mids = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert set(SUBJECTIVE_METRICS) <= mids
    assert set(OBJECTIVE_METRICS) <= mids
    # 主观指标有打分规则
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=6).value == "人工评审":
            assert ws.cell(row=r, column=4).value  # scale_desc


def test_review_sheet_structure():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_review(wb, _mock_rows())
    headers = [ws.cell(row=1, column=c).value for c in range(1, 19)]
    assert "问题（输入）" in headers
    assert "系统答案（输出·待评审）" in headers
    assert "人工评分" in headers
    assert "打分规则" in headers


def test_review_row_count_and_applicability():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_review(wb, _mock_rows())
    n = ws.max_row - 1
    # ANS-001: GEN-01/03/04/06/07/RAG-07/TRACE-02 = 7 (no GEN-08, no RAG-06)
    # CIT-001: 同上 + RAG-06 = 8
    # REF-001: GEN-01/03/04/06/07/RAG-07/TRACE-02 = 7
    # ANS-007: 同上 = 7
    # ANS-010: procedure 类 + GEN-08 = 8
    expected = 7 + 8 + 7 + 7 + 8
    assert n == expected, f"expected {expected} rows, got {n}"


def test_review_refusal_row_shows_reason():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_review(wb, _mock_rows())
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "REF-001":
            assert ws.cell(row=r, column=11).value == "是"
            assert ws.cell(row=r, column=12).value == "false_premise"
            break
    else:
        pytest.fail("REF-001 not found in review sheet")


def test_objective_sheet():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_objective(wb, _mock_rows())
    assert ws.max_row - 1 == 5  # 5 questions
    # 超时题的 run_error 应该在
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "ANS-007":
            assert "timeout" in str(ws.cell(row=r, column=19).value)


def test_summary_sheet_with_data():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_summary(wb, _mock_summary())
    mids = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert "RAG-01" in mids
    assert "GEN-02" in mids
    # RAG-01 值在
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "RAG-01":
            assert ws.cell(row=r, column=3).value == 0.371


def test_summary_sheet_empty():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_summary(wb, {})
    assert ws.max_row >= 1  # 至少 header


def test_load_per_question(tmp_path):
    p = tmp_path / "pq.jsonl"
    rows = _mock_rows()
    p.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows),
                 encoding="utf-8")
    loaded = load_per_question(str(p))
    assert len(loaded) == len(rows)
    assert loaded[0]["question_id"] == "ANS-001"


def test_gen08_only_for_step_types():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_review(wb, _mock_rows())
    gen08_qids = set()
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=13).value == "GEN-08":
            gen08_qids.add(ws.cell(row=r, column=1).value)
    assert gen08_qids == {"ANS-010"}  # 只有 procedure 类型


def test_rag06_only_for_cite_questions():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet_review(wb, _mock_rows())
    rag06_qids = set()
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=13).value == "RAG-06":
            rag06_qids.add(ws.cell(row=r, column=1).value)
    assert rag06_qids == {"CIT-001"}  # 只有 must_cite
